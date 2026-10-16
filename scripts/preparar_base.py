"""Resume a BASE_RL_ORÇ.xlsx num CSV pequeno que o painel importa.

A base original tem centenas de milhares de lançamentos (e um XML de ~500 MB
dentro do .xlsx), o que é pesado demais para abrir no navegador. Este script
soma o Movimento por cenário, mês, empresa, centro de custo, dono e níveis
N1–N4, mantendo os nomes de coluna originais.

Uso:
    pip install openpyxl
    python scripts/preparar_base.py BASE_RL_ORÇ.xlsx public/base.csv
"""
import csv
import sys
from collections import defaultdict

from openpyxl import load_workbook

CHAVES = ['Real / Orçado', 'Mês', 'Empresa', 'Descrição C.C', 'Dono C.C', 'N1', 'N2', 'N3', 'N4']
VALOR = 'Movimento'


def texto(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def main(origem, destino):
    ws = load_workbook(origem, read_only=True, data_only=True).worksheets[0]
    linhas = ws.iter_rows(values_only=True)
    # O cabeçalho não fica na 1ª linha (há um link "Voltar para CAPA" acima).
    for _ in range(20):
        cabecalho = [texto(c) for c in next(linhas)]
        if VALOR in cabecalho and 'N1' in cabecalho:
            break
    else:
        sys.exit(f'Cabeçalho com "{VALOR}" e "N1" não encontrado nas primeiras 20 linhas.')
    faltando = [c for c in CHAVES if c not in cabecalho]
    if faltando:
        sys.exit(f'Colunas não encontradas: {", ".join(faltando)}')
    idx = [cabecalho.index(c) for c in CHAVES]
    i_valor = cabecalho.index(VALOR)

    somas = defaultdict(float)
    lidas = 0
    for linha in linhas:
        v = linha[i_valor] if i_valor < len(linha) else None
        if not isinstance(v, (int, float)) or not v:
            continue
        somas[tuple(texto(linha[i]) for i in idx)] += v
        lidas += 1

    with open(destino, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(CHAVES + [VALOR])
        for chave, soma in sorted(somas.items()):
            if round(soma, 2):
                w.writerow(list(chave) + [f'{soma:.2f}'])
    print(f'{lidas} lançamentos → {len(somas)} linhas resumidas em {destino}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
