"""Concilia o painel com os valores salvos na aba "DRE Gerencial 2026" (filtros TODOS).

Precisa do .xlsb na pasta dre_streamlit/ (ou em DRE_ARQUIVO); sem ele, os testes são pulados.
"""
import sys
from pathlib import Path

import pytest
from openpyxl.utils import column_index_from_string, get_column_letter
from pyxlsb import open_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import config as cfg  # noqa: E402
import dre  # noqa: E402
import leitura  # noqa: E402

pytestmark = pytest.mark.skipif(not cfg.ARQUIVO.exists(), reason=f"{cfg.ARQUIVO.name} não encontrado")
CENTAVO = 0.005  # "diferença 0" = igual até o centavo (o Excel guarda resíduos de ponto flutuante)


@pytest.fixture(scope="module")
def contexto():
    with open_workbook(str(cfg.ARQUIVO)) as wb, wb.get_sheet(cfg.ABA_DRE) as sh:
        excel = {(c.r + 1, c.c): c.v for linha in sh.rows(sparse=True) for c in linha}
    filtros = {r: excel.get((7 + i, column_index_from_string("D") - 1)) for i, r in enumerate(cfg.FILTROS)}
    assert all(v == cfg.TODOS for v in filtros.values()), f"a planilha foi salva com filtros: {filtros}"
    inicio = {}
    for (linha, c), v in excel.items():
        if linha == cfg.LINHA_PERIODOS_DRE and isinstance(v, str):
            nome = next((p for p in cfg.PERIODOS if v.strip().upper().startswith(p)), None)
            if nome:
                inicio[nome] = c
    assert set(inicio) == set(cfg.PERIODOS), f"períodos não achados: {set(cfg.PERIODOS) - set(inicio)}"
    base, estr = leitura.carregar_base(), leitura.carregar_estrutura()
    return {
        "excel": excel, "inicio": inicio, "estr": estr, "res": dre.calcular(base, estr),
        "sinais": leitura.carregar_sinais(), "corte": leitura.mes_de_corte(),
    }


def coluna(ctx, periodo: str, campo: str) -> int:
    return ctx["inicio"][periodo] + cfg.DESLOCAMENTO[campo]


def natureza(ctx, linha):
    return ctx["estr"].set_index("linha").loc[linha, "natureza"]


def excel_num(ctx, linha, c):
    v = ctx["excel"].get((linha, c))
    return v if isinstance(v, (int, float)) else None


@pytest.mark.parametrize("mes", cfg.CONCILIACAO_MESES)
@pytest.mark.parametrize("nome,linha", cfg.CONCILIACAO_LINHAS.items())
def test_criterio_de_aceite(contexto, nome, linha, mes):
    """Critério 1: VENDAS, ROB, ROL, CPV, MC, RO, EBTDA e LL GERENCIAL em jan, fev e mar.

    R 2026 e R 2025 têm de bater sempre. O Δ é comparado quando a planilha tem a fórmula
    (K252, R252 e Y252 estão vazias no Excel; ver test_celulas_vazias_na_planilha).
    """
    periodo = cfg.MESES[mes - 1]
    painel = dre.celulas(contexto["res"], linha, natureza(contexto, linha), periodo,
                         contexto["corte"], contexto["sinais"])
    for campo in ("r26", "r25", "delta"):
        c = coluna(contexto, periodo, campo)
        excel = excel_num(contexto, linha, c)
        if campo == "delta" and excel is None:
            continue
        assert getattr(painel, campo) == pytest.approx(excel, abs=CENTAVO), f"{nome} {get_column_letter(c + 1)}{linha}"


def comparar(ctx, campos):
    """Todas as linhas × todos os períodos. Devolve {célula: motivo} das que não batem."""
    res, ex = ctx["res"], ctx["excel"]
    out = {}
    for _, r in ctx["estr"].iterrows():
        for periodo in cfg.PERIODOS:
            painel = dre.celulas(res, r.linha, r.natureza, periodo, ctx["corte"], ctx["sinais"])
            for campo in campos:
                c = coluna(ctx, periodo, campo)
                ref = f"{get_column_letter(c + 1)}{r.linha}"
                bruto = ex.get((r.linha, c))
                if bruto in (None, "") and r.linha in res.em_branco:
                    continue
                excel = bruto if isinstance(bruto, (int, float)) else None
                tol = CENTAVO if campo in ("r26", "r25", "delta") else 1e-6
                if excel is not None and abs(getattr(painel, campo) - excel) <= tol:
                    continue
                out[ref] = motivo(ctx, r.linha, periodo, campo, excel)
    return out


def motivo(ctx, linha, periodo, campo, excel):
    """Explica uma divergência; None quando não há explicação conhecida."""
    if excel is None:
        return "célula vazia na planilha"
    if campo in ("p26", "p25"):
        valor = excel_num(ctx, linha, coluna(ctx, periodo, "r26" if campo == "p26" else "r25"))
        for base in (13, 14, 16, 51):
            b = excel_num(ctx, base, coluna(ctx, periodo, "r26" if campo == "p26" else "r25"))
            if valor is not None and b and abs(valor / b - excel) < 1e-9:
                return f"planilha divide pela linha {base} (regra: {dre.linha_base(linha)})"
    if campo == "pdelta":
        r25 = excel_num(ctx, linha, coluna(ctx, periodo, "r25"))
        if r25 is not None and abs(r25) < CENTAVO:
            return "R 2025 ≈ 0: a planilha divide por resíduo de ponto flutuante"
    return None


def test_valores_em_todas_as_linhas_e_periodos(contexto):
    """R 2026 e R 2025 batem em todas as linhas e períodos, inclusive ACUMULADO e TOTAL."""
    assert comparar(contexto, ("r26", "r25")) == {}


def test_divergencias_de_delta_e_percentual_sao_da_planilha(contexto):
    """Onde Δ ou % divergem, a causa está na planilha (vazio, outra base ou divisão por ~0)."""
    sem_explicacao = {k: v for k, v in comparar(contexto, ("delta", "p26", "p25", "pdelta")).items() if v is None}
    assert sem_explicacao == {}


def test_celulas_vazias_na_planilha(contexto):
    """Fórmulas que faltam na planilha (itens 3 da análise): só nas linhas 50, 55, 58, 251 e 252."""
    vazias = [k for k, v in comparar(contexto, ("delta", "p26", "p25", "pdelta")).items()
              if v == "célula vazia na planilha"]
    linhas = {int("".join(ch for ch in k if ch.isdigit())) for k in vazias}
    assert linhas <= {50, 55, 58, 251, 252}
